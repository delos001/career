#!/usr/bin/env python3
"""
jd_extract.py - job-description / role-communication text extractor

Takes one job description (or a related communication such as a recruiter
email), pulls the plain text out of it, and prints that text. Used by the
role-intake skill, Phase 1 (JD + comms ingestion).

The input can be a file (.txt, .md, .docx, .pdf) or a web address (http/https).
On success the extracted text is printed to the normal output. On any problem -
an unsupported file type, a missing file, or a file/page that yields no text -
the script prints an error and exits with a non-zero code; it never prints a
half-empty result (this matches the "never proceed on partial content" rule in
rules/global-rules.md).

Author    : Jason Delosh
Created   : 2026-05-14
Project   : career
Usage     : python scripts/ingest/jd_extract.py <path-or-url>
Depends   : python-docx, pypdf, requests, beautifulsoup4
"""

import os
import re
import sys


# ---------------------------------------------------------------------------
# Extractors - one for each kind of source
# Each function below knows how to get the text out of one type of source: a
# plain-text file, a Word document, a PDF, or a web page. They all take the
# source and return its text as a single string. The extract() function further
# down is what decides which of these to call.
# ---------------------------------------------------------------------------

def _from_txt(path):
    """Read a plain-text or markdown file and return everything in it as text.

    The file is read using UTF-8 (the standard text encoding). If a stray byte
    cannot be decoded it is replaced with a placeholder character rather than
    crashing the script.
    """
    with open(path, 'r', encoding='utf-8', errors='replace') as f:
        return f.read()


def _from_docx(path):
    """Pull the text out of a Word (.docx) document.

    Collects every non-empty paragraph, then every table row (the cells of a
    row joined together with ' | '). Tables are included because job
    descriptions often place requirements or details inside a table.
    """
    from docx import Document
    doc = Document(path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = ' | '.join(c for c in cells if c)
            if line:
                parts.append(line)
    return '\n'.join(parts)


def _from_pdf(path):
    """Pull the text out of a PDF file, one page at a time.

    Pages that contain no readable text (for example, a page that is just a
    scanned image) are skipped.
    """
    from pypdf import PdfReader
    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        text = page.extract_text() or ''
        if text.strip():
            parts.append(text)
    return '\n'.join(parts)


def _from_url(url):
    """Download a web page and return only its readable text.

    Fetches the page's HTML, strips out the parts that are never job-description
    content (scripts, styling, site navigation, page header and footer), then
    flattens what is left into plain lines of text with the blank lines removed.
    """
    import requests
    from bs4 import BeautifulSoup
    headers = {'User-Agent': 'Mozilla/5.0 (career role-intake jd_extract)'}
    resp = requests.get(url, headers=headers, timeout=30)
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, 'html.parser')
    # Remove tags that never carry job-description content so their text does
    # not end up mixed into the result.
    for tag in soup(['script', 'style', 'noscript', 'header', 'footer', 'nav']):
        tag.decompose()
    text = soup.get_text(separator='\n')
    # Trim whitespace from each line and drop the empty ones - converting HTML
    # to text leaves a lot of blank lines behind.
    lines = [ln.strip() for ln in text.splitlines()]
    return '\n'.join(ln for ln in lines if ln)


# ---------------------------------------------------------------------------
# Choosing the right extractor
# Looks at the source the caller gave us and sends it to the matching extractor
# above: first deciding web address vs. file, and then, for files, deciding by
# file type.
# ---------------------------------------------------------------------------

def extract(source):
    """Work out what kind of source this is and extract its text accordingly.

    'source' is either a web address or a path to a file on disk. Web addresses
    are downloaded; files are routed by their type (.txt/.md, .docx, .pdf). An
    unknown type, a missing file, or legacy .doc stops the script with a clear
    error.
    """
    # Is this a web address? A source that starts with 'http://' or 'https://'
    # is treated as a URL to download; anything else is treated as a file path.
    # The check is case-insensitive, so 'HTTP://...' also matches.
    if re.match(r'^https?://', source, re.IGNORECASE):
        return _from_url(source)

    if not os.path.exists(source):
        raise FileNotFoundError(f'source not found: {source}')

    # Route by file extension, lower-cased so '.PDF', '.Docx' etc. still match.
    ext = os.path.splitext(source)[1].lower()
    if ext in ('.txt', '.md'):
        return _from_txt(source)
    if ext == '.docx':
        return _from_docx(source)
    if ext == '.pdf':
        return _from_pdf(source)
    if ext == '.doc':
        # Legacy binary Word (.doc) is a different, older format that
        # python-docx cannot read - ask the user to re-save it as .docx.
        raise ValueError('.doc (legacy binary Word) is not supported; '
                         'save the file as .docx and retry')
    raise ValueError(f'unsupported source type: {ext or "(no extension)"}')


# ---------------------------------------------------------------------------
# Command-line entry point
# This block runs when the script is called directly from the command line. It
# checks the argument, runs the extraction, refuses to return an empty result,
# and prints the extracted text.
# ---------------------------------------------------------------------------

if __name__ == '__main__':
    # Expect exactly one argument - the file path or URL to extract from.
    if len(sys.argv) != 2:
        print('Usage: python scripts/ingest/jd_extract.py <path-or-url>',
              file=sys.stderr)
        sys.exit(1)

    try:
        text = extract(sys.argv[1]).strip()
    except Exception as e:
        print(f'Error: {e}', file=sys.stderr)
        sys.exit(1)

    if not text:
        # An empty result is treated as a failure, not a valid answer, so the
        # calling skill does not proceed on nothing.
        print('Error: extraction produced no text', file=sys.stderr)
        sys.exit(1)

    # Force UTF-8 output so accented or special characters in the job
    # description print correctly on Windows.
    sys.stdout.reconfigure(encoding='utf-8')
    print(text)
